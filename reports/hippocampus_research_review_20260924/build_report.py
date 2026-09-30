"""Create the presentation report from verified aggregate evidence."""
from pathlib import Path
import json
import statistics as st
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
D=json.loads((HERE/'evidence/normalized.json').read_text())
R=D['runs']
doc=Document()
BLACK=RGBColor(0,0,0)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
    s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=BLACK
    s.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.size=Pt(12)
for sty in doc.styles:
    for border in list(sty.element.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)

for name,size in [('Heading 1',17),('Heading 2',12),('Heading 3',10.5)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(12)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(8.5)

def page_setup(s,landscape=False):
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11.69 if landscape else 8.27)
    s.page_height=Inches(8.27 if landscape else 11.69)
    s.top_margin=s.bottom_margin=Inches(.62)
    s.left_margin=s.right_margin=Inches(.66)
    s.header_distance=s.footer_distance=Inches(.25)
page_setup(doc.sections[0])
head=doc.sections[0].header.paragraphs[0]
head.add_run('HIPPOCAMPUS SEGMENTATION  |  EXPERIMENT REVIEW').font.size=Pt(8)
foot=doc.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('24 September 2026  •  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
for r in foot.runs:r.font.size=Pt(8)
doc.core_properties.title='Hippocampus segmentation methods and experimental results'
doc.core_properties.subject='Verified results for boundary bands, edge consistency, augmentation and contour functions'

def p(text='',bold=False,style=None):
    x=doc.add_paragraph(style=style);r=x.add_run(text);r.bold=bold
    return x
def h(text,level=1):doc.add_heading(text,level)
def newpage():doc.add_page_break()
def section(landscape=False):page_setup(doc.add_section(WD_SECTION.NEW_PAGE),landscape)
def caption(text):return p(text,style='Caption')
def fmt(x,n=3):return '—' if x is None else f'{x:.{n}f}'
def delta(x,n=3):return '—' if x is None else f'{x:+.{n}f}'
def num(x):return f'{int(x):,}'

def table(headers,rows,widths=None,size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    width=10.37 if doc.sections[-1].orientation==WD_ORIENT.LANDSCAPE else 6.95
    widths=widths or [width/len(headers)]*len(headers)
    scale=width/sum(widths);widths=[x*scale for x in widths]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+name);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders)
    for idx,values in enumerate([headers]+list(rows)):
        row=t.rows[0] if idx==0 else t.add_row()
        rp=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');rp.append(no)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');rp.append(repeat)
        for j,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243E53' if idx==0 else ('F1F5F7' if idx%2==0 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1.04
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            run=para.add_run(str(value));run.font.size=Pt(size);run.font.bold=idx==0
            run.font.color.rgb=RGBColor(255,255,255) if idx==0 else BLACK
    p().paragraph_format.space_after=Pt(1)
    return t

def mt(text):
    x=OxmlElement('m:r')
    rp=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rp.append(sty);x.append(rp)
    e=OxmlElement('m:t');e.text=text;x.append(e);return x

def put(parent,parts):
    if not isinstance(parts,(list,tuple)):parts=[parts]
    for part in parts:parent.append(mt(part) if isinstance(part,str) else deepcopy(part))
    return parent

def script(base,sub=None,sup=None):
    tag='sSubSup' if sub is not None and sup is not None else ('sSub' if sub is not None else 'sSup')
    x=OxmlElement('m:'+tag);put_elem=OxmlElement('m:e');put(put_elem,base);x.append(put_elem)
    for k,val in [('sub',sub),('sup',sup)]:
        if val is not None:x.append(put(OxmlElement('m:'+k),val))
    return x

def frac(a,b):
    x=OxmlElement('m:f')
    for tag,v in [('num',a),('den',b)]:x.append(put(OxmlElement('m:'+tag),v))
    return x

def summation(limits,body):
    x=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc)
    hide=OxmlElement('m:supHide');hide.set(qn('m:val'),'1');pr.append(hide);x.append(pr)
    x.append(put(OxmlElement('m:sub'),limits));x.append(OxmlElement('m:sup'));x.append(put(OxmlElement('m:e'),body))
    return x

def eq(parts):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(5);para.paragraph_format.space_after=Pt(8)
    om=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath');put(math,parts)
    om.append(math);para._p.append(om)

pi=script('p',sub='i');pj=script('p',sub='j');yi=script('y',sub='i');yj=script('y',sub='j')
ei=script('e',sub='i');ej=script('e',sub='j')
Ec=script('E',sup='c');rc=script('r',sup='c');Cv=script('C',sub='valid')
Le=script('L',sub='edge');Lec=script('L',sub='edge',sup='c');Lb=script('L',sub='band');Lbc=script('L',sub='band',sup='c')
Ld=script('L',sub='Dice');lamb=script('λ',sub='band');lame=script('λ',sub='edge')

def selected(budget=75,n=10,fold=None,aug='none',method=None):
    return [r for r in R if r['budget']==budget and r['train_cases']==n and (fold is None or r['fold']==fold) and r['augmentation']==aug and (method is None or r['method']==method)]
def baseline75(seed,fold=0,n=10):return next(r for r in selected(75,n,fold,'none','Baseline') if r['seed']==seed)

p('Hippocampus segmentation methods and experimental results',style='Title')
p('Boundary bands  Edge consistency  Data efficiency  Contour functions',style='Subtitle')
p('Research review prepared for Filippo Focaccia\nEvidence reviewed on 24 September 2026')
h('What the experiments show')
p('The clearest result is faster and better segmentation under the 75-epoch low-data policy. With ten training cases, bands improve on the unaugmented early-stopped Dice baseline; adding uniform edge consistency improves further in every evaluated fold–seed pair. Across folds 0–2 and seeds 0–2, mean validation Dice is 72.633% for baseline, 76.462% for bands and 77.144% for bands plus edge.')
p('The benefit depends on the setting. Longer schedules substantially reduce the bands advantage. Full-data augmentation improves validation while reducing clean-training fit. Contour-function models provide a small surface benefit in a separate protocol, but they also remove true thin structures. None of these findings establishes anatomical correctness or an independent test-set benefit.')
h('How to use this report')
p('The main text gives the high-level explanation first, then the supporting measurements and mathematics. One master table contains all 43 verified runs whose original maximum budget was 75 epochs: the unaugmented baseline and bands across four folds, their augmented fold-0 counterparts, the available larger fractions, and the nine bands-plus-edge runs. The appendix contains the training and validation error distribution for all 61 low-data runs, including longer schedules.')
table(['Presentation segment','Minutes'],[
('Question, methods and comparison rules','0–8'),('75-epoch master table and error pattern','8–22'),('Data fractions and optimization budgets','22–32'),('Augmentation and full-data comparison','32–40'),('Edge loss mathematics','40–48'),('Step contours, learned heads and lessons','48–57'),('Conclusions and discussion','57–60')],[5.7,1.25])
p('Coverage corrections',bold=True)
p('The 10% job was cancelled before starting; no completed 20% run was found. The completed intermediate fractions are 12.5% and 25%. The 200/400/600-epoch experiments use ten cases and do not match 50 full-data epochs in optimizer updates. These are explicit status and budget findings, not missing values silently replaced by another experiment. [S1, S4, S6]')

newpage();h('Methods at a glance')
h('Baseline',2)
p('SwinUNETR predicts background, anterior hippocampus and posterior hippocampus. The low-data and augmentation comparisons use ordinary three-class soft Dice training, with an unaugmented early-stopped run as the reference. The reported validation score is the mean of the anterior and posterior hard Dice per case, then averaged across cases; background is excluded from this reported score. Training DiceLoss includes background. These two definitions must not be interchanged.')
h('Boundary bands',2)
p('Bands tell the model which side of the annotated outer boundary each nearby voxel belongs to. A two-step, six-connected erosion defines an inner foreground band; dilation defines an outer background band. The loss averages foreground binary cross-entropy separately on the two sides, then weights them equally. It adds local foreground/background supervision to Dice. It does not introduce an independent anatomical annotation, and it does not directly supervise the anterior/posterior interface.')
h('Bands plus edge consistency',2)
p('The edge term asks adjacent voxels to reproduce the annotated change in foreground occupancy across their shared face. A face wholly inside or outside foreground should have zero change; a face crossing the boundary should have a signed unit change. This couples neighboring prediction errors. The tested version adds no head, degree weighting or inference-time processing. Section Edge loss mathematics gives the exact equation and its limits.')
h('Augmentation',2)
p('The mild_v1 policy presents transformed versions of the training images and labels while preserving the segmentation task. It changes the training distribution rather than supplying a separate anatomical target. All comparisons here evaluate clean, unaugmented training and validation inputs. It helps at full data and under longer low-data training, but hurts the short low-data baseline; this is evidence of schedule dependence rather than a universally beneficial switch.')
h('Step contours and function heads',2)
p('These methods represent the lower and upper outer boundary as functions along sagittal sections. Early experiments fitted step curves to existing predictions; later heads predicted presence and edge positions, and a differentiable renderer fed them back into the segmentation. The key distinction is between representing a boundary accurately, learning a better boundary from the MRI, and improving the final segmentation. The experiments test these separately.')
h('Comparison rules',2)
p('Every primary gain is anchored to an unaugmented early-stopped baseline with matching fold, seed and data fraction where available. Augmented comparisons retain that anchor and also show the incremental bands effect within the same augmentation policy. Longer runs are compared with their matching unaugmented schedule when it has early stopping, and with the 75-epoch early-stopped reference when necessary. The 600-epoch runs have no early stopping and are explicitly secondary schedule experiments. The contour study has its own early-stopped Dice plus cross-entropy baseline; its scores are not interchangeable with the low-data Dice-only study.')

newpage();h('Experimental setting and error definitions')
p('MSD hippocampus data are split into 208 training and 52 validation case files per fold. The 5% subset uses floor(0.05 × 208) = 10 case files, shared across seeds within a fold. Seeds change initialization and stochastic order, not the subset. Images are padded to 64³ without resizing. The 75-epoch low-data policy uses batch size 1, AdamW learning rate 10⁻⁴ and weight decay 10⁻⁵, StepLR interval 20 with factor 0.5, AMP, and a five-epoch auxiliary ramp. Early stopping uses minimum 60 epochs, patience 8 and minimum improvement 0.0005. The selected checkpoint is the exact maximum logged validation Dice; the stopping decision uses the separate minimum-improvement rule. [S1–S3]')
table(['Error','Exact meaning','Why it matters'],[
('False negative or FN','True anterior/posterior voxel predicted as background','Missed tissue'),
('False positive or FP','True background voxel predicted as anterior/posterior','Extra tissue'),
('A/P swap','True anterior predicted posterior or vice versa','Wrong internal label despite correct foreground occupancy'),
('All class errors','FN + FP + A/P swaps','Mutually exclusive categories; no double counting'),
('Boundary union errors','Inner-band FN + outer-band FP','Local outer-boundary occupancy errors; excludes A/P swaps')],[1.4,3.6,1.95])
p('Error counts are pooled within each split; Dice is a mean across cases. A lower Dice error is therefore not numerically equivalent to fewer pooled wrong voxels. The error percentages in the appendix are shares of all mislabeled voxels, not rates per foreground voxel. Training and validation counts must not be compared directly because their cohort sizes differ. Use errors per case or within-split paired changes.')
p('Two-step morphological bands use six-face connectivity. A two-voxel morphological band is not identical to a Euclidean 2 mm surface band; the contour studies sometimes use physical-distance strata. Those denominators are labeled separately. A/P swaps can lie outside the outer-boundary bands and are counted throughout the foreground in the master table.')
h('Reading the master table',2)
p('N is the number of training cases; F/S is fold/seed; augmentation is No or Mild. Selected/stop gives the chosen checkpoint and actual last epoch. D train and D val are hard foreground macro Dice percentages at the selected checkpoint. Δ base is D val minus the matching unaugmented baseline under the same maximum-75 policy. FN, FP and Swap are whole-volume validation counts on 52 cases. A zero delta identifies the unaugmented baseline. A dash denotes an unavailable experiment, never a zero score.')
p('For folds 0–2 without augmentation, all three models use the latest common audit inference; the remaining rows use the preserved consolidated audit. Small AMP inference differences from older reports are expected and do not represent retraining. Fold 3 has baseline and bands only; no bands-plus-edge run exists there. There are no augmented bands-plus-edge runs in this ledger. [S1–S3]')

section(True);h('All original 75 epoch budget runs')
caption('Table 1  One continuous master table with repeated headers. Scores in percent; Δ base in percentage points. Validation FN and FP include errors beyond the two-step bands. Sources S1–S3.')
rows=[]
for r in R:
    if r['budget']!=75:continue
    v=r['splits']['validation'];t=r['splits']['train_clean']
    rows.append([r['id'],r['train_cases'],f"{r['fold']}/{r['seed']}",'No' if r['augmentation']=='none' else 'Mild',r['method'],f"{r['selected']}/{r['last_logged_epoch']}",fmt(t['dice_pct'],2),fmt(v['dice_pct'],3),delta(r['delta_unaug_baseline_pp']),num(v['fn']),num(v['fp']),num(v['swaps'])])
table(['Run','N','F/S','Aug','Method','Selected\n/ stop','D train','D val','Δ base','FN','FP','Swap'],rows,[.42,.4,.44,.5,1.03,.66,.65,.66,.65,.72,.78,.65],8.5)
p('Unavailable cells are not fabricated: fold-3 bands+edge and all augmented bands+edge models were not run. The 12.5% and 25% entries use 26 and 52 training cases respectively. No completed 10% or 20% result exists in the retrieved evidence.',style='Caption')

section(False);h('What the 75 epoch results mean')
rows=[]
for fold in (0,1,2,3):
    vals={m:[r['splits']['validation']['dice_pct'] for r in selected(75,10,fold,'none',m)] for m in ('Baseline','Bands','Bands + edge')}
    means={m:st.mean(x) if x else None for m,x in vals.items()}
    rows.append([fold,fmt(means['Baseline']),fmt(means['Bands']),fmt(means['Bands + edge']),delta(means['Bands']-means['Baseline']),delta(means['Bands + edge']-means['Baseline']) if means['Bands + edge'] else 'Not run'])
table(['Fold','Baseline','Bands','Bands + edge','Bands Δ base','Edge Δ base'],rows,[.5,1,1,1.25,1.15,1.15],9)
p('Across the three folds where all methods exist, bands add 3.829 percentage points over baseline; bands plus edge add 4.511 points over baseline and 0.682 points over bands. Bands plus edge improves on both controls in all nine fold–seed pairs. Fold 3 supplies additional baseline-versus-bands evidence only. Do not average four folds for bands and three folds for edge and call that a matched contrast.')
h('Error tradeoff against the baseline',2)
rows=[]
for method in ('Baseline','Bands','Bands + edge'):
    rr=[r for r in selected(75,10,None,'none',method) if r['fold'] in (0,1,2)]
    v=[r['splits']['validation'] for r in rr]
    rows.append([method,fmt(st.mean(x['dice_pct'] for x in v)),*[fmt(st.mean(x[k] for x in v),1) for k in ('fn','fp','swaps','boundary_union_errors')]])
table(['Method','D val','All FN','All FP','A/P swaps','Band FN + FP'],rows,[1.6,.85,1.1,1.1,1.1,1.3],8.5)
p('These are average counts per 52-case evaluation over the same nine configurations, not counts from nine independent cohorts. The edge benefit mainly removes extra foreground; it can increase missed foreground relative to bands. That tradeoff matters for thin tissue. The outer-foreground loss has no direct A/P target, so better outer contours must not be reported as proven improvement of the internal anatomical partition.')
h('Does it learn faster',2)
rows=[]
for f in (0,1,2):
    for s in (0,1,2):
        r=next(r for r in selected(75,10,f,'none','Bands + edge') if r['seed']==s)
        b=next(r for r in selected(75,10,f,'none','Bands') if r['seed']==s)
        c=r['learning']['crossings']['bands_reference_best']
        rows.append([f'{f}/{s}',b['selected'],c['first']['first_epoch'] if c['first'] else 'Not reached',c['sustained_three']['confirmed_epoch'] if c['sustained_three'] else 'Not reached',r['last_logged_epoch']])
table(['Fold/seed','Bands best epoch','Edge first reaches it','Three-epoch confirmation','Edge stops'],rows,[.8,1.35,1.6,1.8,1.1],8.5)
p('These are threshold-attainment comparisons on recorded curves, not an equal-quality wall-clock benchmark. Ten updates occur per epoch. Hardware slices and checkpoint I/O differ between some historical runs; calibration also costs compute. The reference target is itself selected on validation, so the crossings are descriptive development evidence. [S2, S3]')

newpage();h('Data fractions and optimizer budgets')
h('Completed fractions and missing runs',2)
rows=[]
for label,n in [('5%',10),('12.5%',26),('25%',52)]:
    a=baseline75(0,n=n);b=next(r for r in selected(75,n,0,'none','Bands') if r['seed']==0)
    rows.append([label,n,fmt(a['splits']['validation']['dice_pct']),fmt(b['splits']['validation']['dice_pct']),delta(b['splits']['validation']['dice_pct']-a['splits']['validation']['dice_pct']),f"{a['last_logged_epoch']}/{b['last_logged_epoch']}"])
table(['Fraction','Cases','Baseline','Bands','Δ baseline','Stop D/B'],rows,[1,.6,1.2,1.2,1.3,1.15])
p('The observed bands gain becomes small as more labeled cases are used: +2.719 points at ten cases for seed 0, +0.211 at 26, and +0.061 at 52. The larger fractions have one seed and one fold only; no bands+edge result exists for them. This is not sufficient evidence for a universal label-scarcity interaction because more cases also mean more updates per epoch. [S1, S4]')
table(['Requested fraction','Verified status'],[
('10% or 20 training cases','Job 662625 was staged and submitted, then cancelled before starting. No trained result.'),
('20%','No completed run found in the searched local ledger and cluster fraction roots.'),
('12.5% or 26 cases','Completed; included because this is the actual intermediate fraction in the audited record.'),
('25% or 52 cases','Completed; fold 0, seed 0, baseline and bands.')],[1.5,5.45])
h('Exact update accounting',2)
eq(['U(E,N,b) = E × ⌈N/b⌉,       U(50,208,1) = 10,400'])
p('This formula assumes one training pass per epoch, no gradient accumulation and no skipped optimizer steps. At batch size one, each case produces one update. For the 600-epoch experiment the protocol additionally checked AMP skips. Calibration steps are overhead and are not included in the model-training update totals.')
table(['Training cases','Epochs needed for 10,400 updates','Evidence here'],[
('10 or approximately 5%','1,040','200/400/600 schedules use these ten cases'),
('20 or approximately 10%','520','Cancelled before training'),
('26 or 12.5%','400','Only maximum-75 runs completed here'),
('41 or approximately 20%','254 gives 10,414 updates','No completed run found'),
('52 or 25%','200','Only maximum-75 runs completed here')],[1.7,2.35,2.9],8.5)
p('Therefore the existing 200-, 400- and 600-epoch runs cannot be described as equal-update matches to 50 full-data epochs: they use the ten-case subset, not the corresponding larger fractions. Full-data early stopping actually ended at 25/29/29 epochs for baseline seeds 0/1/2: 5,200/6,032/6,032 updates. The 600-epoch ten-case run is close to a 29-epoch stopped baseline, but not equal to the nominal 50-epoch budget. Equal counts would still not match label diversity or learning-rate exposure. [S5, S6]')

section(True);h('Longer schedules and the early stopped reference')
caption('Table 2  Every separate 200/400/600-budget run. Selected and final Dice are distinct. Δ75ES compares the selected score with the same-seed unaugmented 75-epoch early-stopped baseline on the same ten cases. Within-policy Δ isolates bands versus Dice for the same schedule and augmentation; it is secondary when early stopping is disabled.')
rows=[]
for r in R:
    if r['budget']==75:continue
    ref=baseline75(r['seed']);matched=next(x for x in R if all(x[k]==r[k] for k in ('budget','seed','fold','augmentation','train_cases')) and x['method']=='Baseline')
    v=r['splits']['validation']['dice_pct'];final=r['final_splits']['validation']['dice_pct']
    rows.append([r['id'],r['budget'],r['seed'],'No' if r['augmentation']=='none' else 'Mild',r['method'],'Yes' if r['early_stopping_patience'] else 'No',f"{r['selected']}/{r['last_logged_epoch']}",num(r['optimizer_updates_at_stop']),fmt(v),fmt(final),delta(v-ref['splits']['validation']['dice_pct']),delta(v-matched['splits']['validation']['dice_pct'])])
table(['Run','Budget','Seed','Aug','Method','ES','Selected\n/ stop','Updates\nat stop','Best D val','Final D val','Δ75ES','Within-policy Δ'],rows,[.4,.55,.4,.5,.85,.4,.75,.85,.85,.85,.7,.95],8.5)
p('The 200-budget nonaugmented runs stop at 160/160/169 epochs for baseline and 160/160/160 for bands. The 400-budget seed-2 pair stops at 320. The 200-budget augmented runs and all 600-budget runs disable early stopping; selecting their best checkpoint afterward does not make them early-stopped runs. [S1, S6]')
section(False);h('Longer training and generalization')
h('What changes with more optimization',2)
p('At the nonaugmented 200-epoch maximum, mean selected Dice is 79.285% for baseline and 80.024% for bands: +0.739 points, versus +3.935 at maximum 75 on fold 0. The separate seed-2 400-budget pair gives 79.519% versus 79.874% (+0.356). This supports an early optimization advantage for bands; it does not prove convergence.')
p('The 600-epoch fits use StepLR interval 160, versus 53 for separate 200-budget fits and 106 for 400-budget fits. Duration and learning-rate timing therefore change together. Best-checkpoint gains can disappear at the final epoch, as the next table shows.')
h('Snapshots within the same 600 epoch schedule',2)
rows=[]
for aug in ('none','mild_v1'):
    rr=selected(600,10,0,aug)
    a=next(r for r in rr if r['method']=='Baseline');b=next(r for r in rr if r['method']=='Bands')
    for epoch in (200,400,600):
        key='latest' if epoch==600 else f'epoch{epoch}'
        av=next(c for c in a['checkpoints'] if c['key']==key)['summary'];bv=next(c for c in b['checkpoints'] if c['key']==key)['summary']
        rows.append(['No' if aug=='none' else 'Mild',epoch,num(epoch*10),fmt(av['validation']['dice_pct']),fmt(bv['validation']['dice_pct']),delta(bv['validation']['dice_pct']-av['validation']['dice_pct'])])
table(['Aug','Epoch','Updates','Baseline D val','Bands D val','Bands Δ'],rows,[.6,.7,1,1.6,1.6,1.2],9)
p('These snapshots are not interchangeable with separately trained 200- or 400-budget runs because the learning-rate schedules differ. All six rows refer to the same seed-2 600-epoch schedule. They show the bands gap shrinking under extended optimization, especially with augmentation. No bands+edge long-schedule model was trained, so convergence evidence for bands must not be attributed to edge consistency.')
h('Learning versus generalization',2)
p('For the nonaugmented 600-epoch pair, clean-training Dice reaches 98.927% baseline and 99.682% bands while validation remains around 79%. Between epochs 501–550 and 551–600, mean validation changes by only +0.004 and −0.020 percentage points, while supervised losses still decline by 7.12% and 10.82%. Validation has plateaued even though fitting continues.')
p('For the augmented pair, final clean-training Dice is 98.130% and 98.570%; validation is approximately 86%. Over the same late windows, supervised losses change by less than 1%, and validation changes by +0.036 and +0.004 points. This is stronger descriptive evidence of a joint plateau, but still a single seed and a reused development fold. Reaching 600 epochs is not itself proof of convergence. [S1, S6]')
h('Comparison with the full data reference',2)
p('The full-data, unaugmented, early-stopped seed-2 baseline selects epoch 21 and scores 87.895% validation Dice. The ten-case augmented 600-epoch baseline scores 85.866%; augmented bands scores 86.183% at its selected checkpoint. They remain below the full-data reference by 2.029 and 1.712 points respectively. These are context comparisons across different label counts and schedules, not matched causal estimates of augmentation or bands.')

newpage();h('Early stopped augmentation versus no augmentation')
p('The full-data replication is the cleanest comparison for augmentation: 208 training and 52 validation cases, fold 0, three independent model initializations, the same Dice objective and optimizer, maximum 50 epochs, minimum 25, patience 8, delta 0.0005. Each method chooses its own validation-best checkpoint before evaluation on clean inputs. [S5]')
rows=[]
for x in D['full_aug_comparison']:
    if x['checkpoint']!='best':continue
    last=next(z for z in D['full_aug_comparison'] if z['seed']==x['seed'] and z['checkpoint']=='latest')
    rows.append([x['seed'],f"{x['baseline_epoch']}/{last['baseline_epoch']}",f"{x['augmentation_epoch']}/{last['augmentation_epoch']}",fmt(x['baseline_train_dice_pct']),fmt(x['augmentation_train_dice_pct']),fmt(x['baseline_validation_dice_pct']),fmt(x['augmentation_validation_dice_pct']),delta(x['validation_dice_delta_pp'])])
table(['Seed','Base sel/stop','Aug sel/stop','Base train','Aug train','Base val','Aug val','Δ base'],rows,[.45,.95,.95,1,1,1,1,.85],8.5)
p('Mean validation Dice increases from 87.839% to 88.642%, or +0.802 points. All three seeds improve. Clean-training Dice decreases under augmentation, which is compatible with stronger regularization and a more difficult training distribution. It does not by itself prove a causal regularization mechanism.')
rows=[]
for x in D['full_aug_comparison']:
    if x['checkpoint']!='best':continue
    for method,prefix in [('Baseline','baseline'),('Augmented','augmentation')]:
        rows.append([x['seed'],method,num(x[prefix+'_validation_fn']),num(x[prefix+'_validation_fp']),num(x[prefix+'_validation_swaps']),num(x[prefix+'_validation_wrong'])])
table(['Seed','Method','Validation FN','Validation FP','A/P swaps','All errors'],rows,[.45,1.3,1.3,1.3,1.2,1.3],8.5)
p('Augmentation reduces total validation errors in every seed, but A/P swaps are not consistently reduced. The main benefit is foreground segmentation. Full-data training/validation error counts for all six selected models are retained in the appendix.')
h('Why the low data augmentation result differs',2)
for budget in (75,200):
    base=selected(budget,10,0,'none','Baseline');aug=selected(budget,10,0,'mild_v1','Baseline')
    mean_b=st.mean(r['splits']['validation']['dice_pct'] for r in base);mean_a=st.mean(r['splits']['validation']['dice_pct'] for r in aug)
    p(f'At maximum {budget} epochs on ten cases, baseline validation Dice changes from {mean_b:.3f}% without augmentation to {mean_a:.3f}% with augmentation ({mean_a-mean_b:+.3f} points). '+('Both short policies have early stopping.' if budget==75 else 'Here the nonaugmented runs early-stop, whereas the augmented runs use the full 200 epochs; this is not a pure augmentation contrast.'))
p('The short-policy loss under augmentation and long-policy recovery argue for reporting augmentation jointly with the optimization budget. Do not use the stronger full-data augmentation result as the baseline for claiming a bands-only gain over unaugmented training.')

newpage();h('Full data bands and edge as a context check')
p('The full-data edge experiment used mild augmentation. Its direct Dice comparator was the augmented early-stopped seed-0 model, not the unaugmented baseline. The table includes the latter explicitly, as requested, but differences to it combine augmentation and auxiliary-loss changes. The direct loss comparison is among the three augmented rows. [S5, S7]')
base=next(x for x in D['full_aug_comparison'] if x['seed']==0 and x['checkpoint']=='best')
rows=[['Unaugmented baseline',fmt(base['baseline_train_dice_pct']),fmt(base['baseline_validation_dice_pct']),'0.000',num(base['baseline_validation_fn']),num(base['baseline_validation_fp']),num(base['baseline_validation_swaps'])]]
for name,label in [('dice','Augmented Dice'),('bands','Augmented bands'),('edge','Augmented bands + edge')]:
    x=D['full_edge'][name];v=x['validation']
    rows.append([label,fmt(x['train']['dice_pct']),fmt(v['dice_pct']),delta(v['dice_pct']-base['baseline_validation_dice_pct']),num(v['fn']),num(v['fp']),num(v['swaps'])])
table(['Method','Train Dice','Val Dice','Δ unaug base','Val FN','Val FP','A/P swaps'],rows,[1.6,.8,.8,.95,.9,.9,.8],8.2)
p('Against augmented bands, edge improves training Dice by +0.713 points but changes validation Dice by −0.0118 points. It removes 5,625 training boundary errors but only 105 validation boundary errors; the validation balanced boundary error slightly worsens. Against augmented Dice, validation boundary errors increase by 287. Thus the full-data audit shows improved fitting with limited transfer, whereas the low-data 75-epoch experiment shows gains on both splits. This difference has not been isolated in a factorial experiment, so do not claim that scarcity itself causes the edge benefit.')
h('Full data training error distribution',2)
rows=[]
for name,label in [('dice','Augmented Dice'),('bands','Augmented bands'),('edge','Augmented bands + edge')]:
    t=D['full_edge'][name]['train']
    rows.append([label,num(t['fn']),num(t['fp']),num(t['swaps']),num(t['boundary_union_errors'])])
table(['Method','Train FN','Train FP','Train swaps','Band FN + FP'],rows,[2,1.2,1.2,1.2,1.35],8.5)
p('All rows use the same 208 clean training cases. The unaugmented seed-0 baseline has 33,739 FN, 51,199 FP and 2,832 swaps in the matched augmentation audit; its complete train/validation results are in the appendix.',style='Caption')
h('What anatomical information the losses contain',2)
p('Bands encode the location and side of the annotated foreground boundary. Edge consistency additionally encodes relative occupancy across six-connected voxel faces. Both targets are derived from the existing segmentation label. They supply a useful spatial inductive bias, not a new independent label source. Neither term distinguishes anterior from posterior once they are merged into foreground, nor guarantees one connected component, absence of holes, a particular genus or a valid internal cut.')

newpage();h('Edge loss mathematics')
h('Notation and support',2)
p('For a voxel i, let zᵢ be its three logits and sᵢ = softmax(zᵢ). Define foreground probability pᵢ = sᵢ,A + sᵢ,P and binary target yᵢ = 1 when the label is anterior or posterior. Let F be the annotated foreground. Let erosion and dilation use six-face adjacency, repeated twice.')
eq(['I = F ∖ ',script('erosion',sub='6',sup='2'),'(F),    O = ',script('dilation',sub='6',sup='2'),'(F) ∖ F,    B = I ∪ O'])
p('I is the inner band, O the outer band, and B their union. Within each case c, Eᶜ contains every face-adjacent pair (i,j) for which both endpoints are in B. The implementation enumerates positive x, y and z directions once. It includes same-label faces as well as foreground/background cuts; it does not restrict the loss to crossing faces. Boundary labels build the training support, but are not needed during inference.')
h('Bands term',2)
eq([Lbc,' = ',frac('1','2'),' [ ',frac(summation('i∈I',['−log(',pi,')']),'|I|'),' + ',frac(summation('i∈O',['−log(1−',pi,')']),'|O|'),' ]'])
p('This is the tested zero-focal-gamma, unweighted formulation. Each side is averaged separately, so the larger outer band does not dominate simply by containing more voxels. The implementation uses grouped foreground log-odds for stable binary cross-entropy. Degree weighting is absent from the runs in this report.')
h('Signed face consistency',2)
eq([Lec,' = ',frac('1',['|',Ec,'|']),summation(['(i,j)∈',Ec],script(['[(',pi,'−',pj,')−(',yi,'−',yj,')]'],sup='2'))])
eq([Le,' = ',frac('1',['|',Cv,'|']),summation(['c∈',Cv],Lec)])
p('Cases with no valid face are excluded from the case mean; if none are valid, the implementation returns zero. Otherwise every valid case has equal weight, and every valid face within that case has equal weight. Reversing an edge changes the residual sign but not its squared loss.')
table(['Adjacent targets','Desired probability change','Meaning'],[
('yᵢ = yⱼ','pᵢ − pⱼ = 0','Agreement inside the same anatomical region'),
('yᵢ = 1 and yⱼ = 0','pᵢ − pⱼ = 1','Foreground-to-background boundary'),
('yᵢ = 0 and yⱼ = 1','pᵢ − pⱼ = −1','Same boundary in reverse orientation')],[1.65,1.8,3.5],9)
p('For a crossing face with pᵢ = 0.8 and pⱼ = 0.2, the face penalty is [(0.8−0.2)−1]² = 0.16. It asks for a sharper correct jump. For a same-label face with 0.8 and 0.6, the penalty is 0.04. It asks for matching probabilities there. Ordinary probability smoothing would also penalize the correct crossing jump; this target-subtracted term does not.')

newpage();h('Why the edge term can complement bands')
p('Write each voxel probability error as eᵢ = pᵢ − yᵢ. Then each face contributes (eᵢ−eⱼ)². If D is an oriented incidence matrix of the band graph, the loss is a quadratic energy in the differences of prediction errors:')
eq([Lec,' = ',frac(script('‖D(p−y)‖',sup='2'),['|',Ec,'|']),' = ',frac([script('(p−y)',sup='T'),script('D',sup='T'),'D(p−y)'],['|',Ec,'|'])])
eq([frac(['∂',Lec],['∂',pi]),' = ',frac('2',['|',Ec,'|']),summation('j∈N(i)',['(',ei,'−',ej,')'])])
p('The band BCE gives a voxelwise direction toward the correct occupancy. The edge term couples that direction to neighboring errors. It is a supervised graph-Laplacian penalty on the error field. The derivative above is with respect to probabilities; the network receives it through softmax and its parameters, so this expression alone does not guarantee a helpful parameter update.')
p('The edge term alone cannot fix a constant error offset on a connected component, because D annihilates constants. Dice and bands anchor the absolute foreground assignment. It also cannot explicitly recover an A/P swap: if probability mass moves from anterior to posterior while their sum stays fixed, p is unchanged. Its success is therefore judged by segmentation and boundary outcomes, not asserted from the graph interpretation.')
h('Total objective and calibration',2)
eq([script('L',sub='total'),'(t) = ',Ld,' + α(t)[',lamb,Lb,' + ',lame,Le,']'])
eq(['α(t) = min(1, t/5)'])
p('The edge coefficient is calibrated separately for each fold and seed using the retained epoch-five Dice checkpoint and that fold’s ten training cases. The bands coefficient is inherited from the completed bands run. At each calibration case, compute RMS gradients with respect to logits for Lbase = LDice + λband Lband and the unweighted edge loss, then form the ratio rᶜ = RMS(∂Ledge/∂z) / RMS(∂Lbase/∂z).')
eq([lame,' = min [ ',frac('0.1',['median(',rc,')']),', ',frac('0.5',[script('Q',sub='0.95'),'(',rc,')']),' ]'])
p('This targets a median edge-gradient ratio of at most 0.1 with a 95th-percentile cap of 0.5. It uses training labels only. These are scale-setting rules, not guarantees of a particular gradient ratio throughout training, and not validation tuning. Fold-0 coefficients range from 0.00761 to 0.00909. After calibration every model trains from scratch; it does not continue the calibration checkpoint. [S2, S3, S8]')
h('Practical implementation',2)
p('Compute softmax in float32, sum anterior and posterior probabilities, subtract the binary foreground mask, and take squared adjacent differences of that residual. Mask faces only when both endpoints are inside B. Sum over the three axes, divide by face count per case, then average valid cases. No wraparound, diagonal edges, learned head or change to inference is used.')

newpage();h('From step contours to boundary prediction')
h('The anatomical observation',2)
p('In the axis-aligned MSD volumes, upper and lower sagittal contours generally descend from posterior to anterior. Across 260 reference masks, the upper-contour median Spearman correlation with anterior position is −0.995 and median decreasing step-fit residual is 0.141 voxel. However, the early-stopped baseline already exhibits the pattern: median correlation −0.996 and residual 0.135. A global monotonicity violation is therefore a weak error detector, and 33/260 reference masks contain valid upward jumps exceeding one voxel. A hard universal monotonicity rule would contradict some labels. [S9]')
h('Step fitting and its first gradient screen',2)
p('At fixed sagittal position x, define lower and upper heights Lₓ(y) and Uₓ(y). For probabilities along a superior–inferior ray, soft first-hit mass is f(z) = p(z)∏ₜ<z[1−p(t)]; last-hit mass uses the reverse ray. Conditional expected hit positions yield soft endpoints. Fit decreasing isotonic functions separately within each contiguous occupied y run, and compare predicted and reference fitted heights. Gradients pass through each fixed plateau mean during backward propagation.')
eq(['Lfit = mean [ |I↓(L̂) − I↓(L*)|/Z + |I↓(Û) − I↓(U*)|/Z ] / 2'])
p('An equal-gradient-RMS frozen-logit update on 52 fold-0 cases improved union Dice by only +0.000114 over Dice-only, with 26 cases improving, 4 worsening and 22 tied. A transition term gave +0.000108; a residual term gave +0.000104. These are tiny label-informed local update screens, not trained-network generalization results. The residual term addresses the fact that different wrong raw curves can project to the same isotonic plateau, but did not improve this screen. [S10]')
h('A good representation is not yet a good predictor',2)
p('The two-edge representation nearly reconstructs the labels when given true presence and true endpoints: oracle union Dice is 0.9994. The challenge is learning those quantities accurately from images and converting them into selective corrections. The union contour cannot determine A/P identity, and a sagittal silhouette does not fully constrain lateral shape.')
table(['Experiment','Baseline union Dice','Head or candidate','What it establishes'],[
('Withdrawn pilot A to D','89.27%','89.29% segmentation','Tiny change; interval includes zero'),
('Frozen independent head','89.636%','89.563% head mask','Slightly better conditional edges; worse presence recall'),
('Oracle endpoints and presence','Not a trained model','99.94% reconstruction','Representation ceiling, not deployable performance')],[1.7,1.25,1.25,2.75],8.5)
p('These rows have different baselines and selection protocols; they are not a ranking on one common experiment. The pilot was withdrawn and its heavy artifacts deleted. Its retained summary is historical evidence only. [S11, S12]')

newpage();h('What the learned boundary heads added')
h('Independent head on a frozen early stopped baseline',2)
p('A small head receives frozen decoder features and the normalized MRI, predicting column presence and categorical lower/upper endpoint positions. On the same 4,029 occupied columns where both methods detect foreground but disagree on edges, the head wins 53.6% of non-ties, with a case-grouped interval of 51.2–56.1%. This supports a small complementary edge signal. Yet on 687 presence disagreements, it wins only 43.4%, with interval 38.7–48.4%. It removes 258 false-positive columns but adds 349 false-negative columns. Better conditional edge accuracy does not improve its standalone mask. [S12]')
h('Withdrawn four arm contour pilot',2)
p('The pilot coupled head losses through shared decoder features but did not implement direct training-time enforcement of segmentation logits by the predicted contour. Its later hard inference correction broke more voxels than it fixed: arm C fixed 4,499 and broke 6,388; arm D fixed 4,105 and broke 5,640. For D, the head and segmentation endpoint MAEs were virtually identical, while the head had more false-positive occupied columns. This was a largely redundant estimate with a presence penalty, not a successful contour correction. [S11]')
h('Differentiable function feedback',2)
p('The later function model explicitly connects a residual endpoint/presence head to the voxel output. First/last-hit probabilities anchor predicted endpoints to the raw foreground. A fitted pair of functions defines soft interval occupancy, which modifies the background logit while leaving anterior/posterior log-odds unchanged. Training includes rendered and raw segmentation losses, supervised curves, a ground-truth-anchored boundary voxel term, and an ordering penalty. This is a different model from the head-only pilot. [S13]')
eq(['oᵣ(z) = π̂ᵣ σ[(z−L̂ᵣ+½)/τ] σ[(Ûᵣ−z+½)/τ]'])
eq(['z′BG = zBG − β M tanh[logit(o)/M],     z′A = zA,     z′P = zP'])
eq(['Lseg = ',frac('LDiceCE(z′) + η LDiceCE(z)','1+η')])
p('Presence π̂ says whether the ray exists; L̂ and Û locate its interval. β is bounded and warmed up. The normalized raw auxiliary loss keeps the voxel network useful. The tested function study uses a Dice-plus-cross-entropy baseline, batch size 2, 166 fitting and 42 inner-selection cases within each 208-case training fold, and 52 outer cases. Maximum 50 epochs, minimum 20 and patience 8 select an early-stopped matched baseline. These outcomes cannot be substituted into the Dice-only low-data table.')

newpage();h('Rendered contour results against the matched baseline')
p('Across folds 1–4, the function study compared A, the matched unaugmented early-stopped DiceCE baseline; D, the supervised function/voxel model without training through rendering; D′, D rendered only at inference; and E, trained through the renderer. The baseline comparison is primary here. E versus D′ is a secondary mechanism contrast. One backbone seed was used. [S14]')
rows=[]
for fold,x in D['contour']['fold_summaries'].items():
    a=x['A'];e=x['E'];rows.append([fold,fmt(100*a['mean_union_dice']),fmt(100*e['mean_union_dice']),delta(100*(e['mean_union_dice']-a['mean_union_dice'])),fmt(a['mean_assd_mm'],4),fmt(e['mean_assd_mm'],4)])
table(['Fold','A union Dice','E union Dice','E minus A pp','A ASSD mm','E ASSD mm'],rows,[.5,1.25,1.25,1.25,1.3,1.3],8.5)
rows=[]
for name,x in D['contour']['pooled'].items():
    rows.append([name,fmt(100*x['mean_union_dice']),fmt(x['mean_assd_mm'],4),delta(100*(x['mean_union_dice']-D['contour']['pooled']['A']['mean_union_dice']))])
table(['Model','Union Dice percent','ASSD mm','Dice Δ baseline pp'],rows,[1.1,1.95,1.6,2.3])
p('E improves pooled union Dice by +0.242 points and mean surface distance by approximately −0.0126 mm versus A. Union Dice is slightly lower on folds 1 and 2 and higher on folds 3 and 4; this is less uniform than the low-data edge result. A later common audit reproduces E ASSD as 0.483148 rather than the historical 0.483144; the small discrepancy is documented, not hidden. The following error table uses that common audit. [S14, S15]')
table(['Model','Foreground FN','Foreground FP','A/P swaps','Interpretation'],[
('A baseline','71,949','87,880','17,476','Matched early-stopped reference'),
('E raw','49,281','125,859','18,390','Expansive foreground before rendering'),
('E rendered','73,988','81,778','17,637','Less FP but more missed foreground than A')],[1.1,1,1,1,2.85],8.5)
p('These are pooled counts across 208 outer cases, not the 52-case counts of a single fold. Rendering removes 44,467 FP voxels but deletes 24,961 true foreground voxels, of which 24,181 were correctly labeled before rendering. It changes no anterior label directly into posterior or vice versa; class odds are preserved. Net E versus A reduces foreground/background errors by 4,063, but increases A/P swaps by 161. [S15]')

newpage();h('Why contour feedback remains difficult')
h('Thin foreground and presence',2)
p('In the same common audit, correctly overlapping thin-ray recall falls from 63.12% for A to 55.55% for E, although E raw reaches 83.12%. Thin voxel recall falls from 59.47% to 53.84%. Mean false-positive rays fall from 40.05 to 28.94 per case. E therefore improves average surface accuracy partly through stronger suppression, paying for it with thin-tissue loss. Deep foreground is already recovered: A misses only 41 of 65,015 reference voxels more than 2 mm inside the surface. [S15, S16]')
p('Oracle presence and endpoint interventions show that some mistakes are representationally recoverable, but use labels and are not deployable. Setting true-thin presence to one raises overlap recall to 92.05% while adding false positives and leaving ASSD essentially unchanged. A minimal sufficient oracle presence increase gives ASSD 0.473279 mm and 92.05% overlap recall with only approximately 4.86 extra FP voxels per case. This separates existence prediction from the amount of voxel correction required. [S16, S17]')
h('A favorable output gradient can still fail through shared parameters',2)
p('A diagnostic on fitting batches found that the existing rendered model already pushes erased-thin presence upward in 95.98% of sampled ray observations and pushes all 734 sampled erased true voxels toward foreground. Yet its shared-head SGD direction lowers mean erased-thin presence in 8/16 batches. Increasing focal rescue raises thin and empty-ray presence together. These are local derivative findings, not a reconstructed AdamW training trajectory. They warn against assuming that a larger auxiliary weight gives a selective anatomical correction. [S18]')
h('What can be presented as a supported conclusion',2)
p('A compact lower/upper contour can represent the mask very well, and a head can learn useful edge information. The missing ingredient is reliable, selective correction: a presence or renderer decision must recover true boundary voxels without adding false positives or erasing thin anatomy. The rendered model has a small internal surface benefit, but its error tradeoff and limited seed coverage prevent a broad claim of better anatomical segmentation.')
p('For the practical low-data segmentation result, the simpler bands-plus-edge objective currently has the clearest replicated benefit under the tested 75-epoch policy. It does not replace the contour work; it answers a narrower question with fewer modeling components. Longer-schedule edge experiments and an untouched evaluation cohort are still needed to assess persistence and generalization.')

newpage();h('Conclusions for the presentation')
p('1  Under the maximum-75 low-data policy, bands and bands plus edge both improve on the unaugmented early-stopped baseline. The edge increment repeats across all nine tested fold–seed pairs, with mean Dice 77.144% versus 72.633% for baseline on the same three folds.',bold=True)
p('2  Error distributions matter. Foreground FP reduction explains much of the gain, while FN can rise and A/P swaps are not consistently reduced. The appendix retains every run rather than reporting only winning seeds.')
p('3  Longer training reduces the bands gap. The 200/400/600 runs change optimization duration and learning-rate timing, and some disable early stopping. They are not equal-update matches to the nominal 50-epoch full-data baseline.')
p('4  Augmentation improves the full-data early-stopped baseline by +0.802 points on average across three seeds, but can hurt when ten-case training is too short. This reconciles the apparently conflicting augmentation outcomes.')
p('5  Step functions capture a real contour pattern, but that pattern is already present in predictions. Learned contour feedback must improve patient-specific boundary placement and preserve thin foreground, rather than merely enforce a global shape tendency.')
h('Questions the evidence does not yet answer',2)
p('Does the edge improvement remain after longer low-data schedules? Does it replicate on a genuinely untouched cohort and other datasets? Can a contour head improve thin-versus-empty discrimination under fixed FP constraints? Would a fully update-matched data-fraction experiment change the ranking? None of these questions is settled by the present selected validation checkpoints.')
h('Presentation cautions',2)
p('Say “validation-selected development result,” not “independent test accuracy.” Say “maximum 75 epochs with early stopping,” not “every model ran 75 epochs.” Say “case files,” unless subject/hemisphere identities have been separately verified. Report class macro Dice and union Dice under distinct names. Do not describe an oracle intervention as a learned result, or a no-early-stopping run as early stopped because it has a best checkpoint.')

section(True);h('Appendix training and validation error distributions')
caption('Table 3  Selected-checkpoint errors for every low-data run in the report. Run IDs link to Tables 1 and 2. FP FN AP percentages partition validation mislabeled voxels and sum to 100 apart from rounding. They are not voxel-level error rates. Clean training uses N cases and validation uses 52. S1–S3.')
rows=[]
for r in R:
    tr=r['splits']['train_clean'];v=r['splits']['validation'];sh=v['error_shares_pct']
    rows.append([r['id'],f"{r['budget']} / {r['train_cases']} / {r['fold']}/{r['seed']}",r['method']+(' aug' if r['augmentation']!='none' else ''),num(tr['fn']),num(tr['fp']),num(tr['swaps']),num(v['fn']),num(v['fp']),num(v['swaps']),f"{sh['fp']:.1f} / {sh['fn']:.1f} / {sh['swaps']:.1f}",fmt(v['errors_per_case'],1)])
table(['Run','Budget / N / F/S','Method','Train FN','Train FP','Train AP','Val FN','Val FP','Val AP','Val shares FP / FN / AP %','Val errors / case'],rows,[.4,1,1.25,.7,.8,.7,.8,.8,.7,1.3,.8],8.2)
p('These categories count each wrong voxel once. A/P means swaps, not all foreground errors. Model-to-model comparisons are matched within a fold, seed and data fraction; raw counts from different cohorts should not be subtracted as if they were paired. For the nine edge comparisons, validation and training counts are from common re-inference. Historical audit values are retained for the other runs.',style='Caption')

newpage();h('Appendix spatial location of boundary errors')
caption('Table 4  Validation spatial strata at selected checkpoints for the 37 low-data runs with retained band audits. Inner FN plus core FN equals all FN; outer FP plus far FP equals all FP. All bands use two iterations of six-face morphology. The other 24 runs have categorical errors in Table 3 but no retained spatial split in the normalized evidence. S1–S3.')
rows=[]
for r in R:
    v=r['splits']['validation']
    if 'inner_band_fn' not in v:continue
    rows.append([r['id'],r['method'],num(v['inner_band_fn']),num(v['core_fn']),num(v['outer_band_fp']),num(v['far_fp']),fmt(100*(v['inner_band_fn']+v['outer_band_fp'])/(v['fn']+v['fp']),2)])
table(['Run','Method','Inner band FN','Deeper core FN','Outer band FP','Far FP','Union errors inside bands %'],rows,[.5,1.45,1.4,1.4,1.4,1.0,1.8],8.5)
p('A/P swaps are counted separately in Table 3 because changing anterior to posterior does not change foreground occupancy. A high concentration of mistakes within these bands explains why local boundary objectives are plausible; it does not make individual near-boundary errors anatomically harmless.',style='Caption')

section(False);h('Appendix full data error distribution')
caption('Table 5  Selected checkpoints in the matched full-data augmentation study. All evaluations use unaugmented inputs. Training has 208 cases and validation 52. S5.')
rows=[]
for x in D['full_aug_comparison']:
    if x['checkpoint']!='best':continue
    for method,prefix in [('Baseline','baseline'),('Augmented','augmentation')]:
        for split in ('train','validation'):
            rows.append([x['seed'],method,'Train' if split=='train' else 'Val',num(x[f'{prefix}_{split}_fn']),num(x[f'{prefix}_{split}_fp']),num(x[f'{prefix}_{split}_swaps']),fmt(x[f'{prefix}_{split}_dice_pct'])])
table(['Seed','Method','Split','FN','FP','A/P swaps','Dice %'],rows,[.4,1.2,.6,1,1,1,1],8.5)
h('Data completeness and provenance',2)
p('The 61-run low-data ledger contains 43 original maximum-75 runs and 18 separate longer-schedule runs. The 43 include four 12.5%/25% pilot runs, six augmented fold-0 runs, 24 unaugmented baseline/bands runs across four folds, and nine unaugmented edge runs across three folds. Error counts and per-run deltas were computed programmatically from preserved JSON summaries; no new training, inference, checkpoint selection or deletion was performed for this document.')
p('For the latest edge audit, actual training stops are fold 1 seeds 0/1/2: 75/75/60; fold 2: 75/60/67. All runs and both split audits completed in job 668085, exit 0, elapsed 33m42s. The retained fold-0 edge job was 667831. The prior degree-weighted A experiment was deleted at the user’s request and is not reconstructed in this report.')
p('Older consolidated results can differ slightly from newer shared AMP inference. The master table consistently uses the newer shared audit for fold-0–2 unaugmented comparisons and identifies other rows as historical. The full-data augmentation and function-head studies have distinct objectives, case-selection schemes and checkpoints, so their baseline anchors are stated separately.')

newpage();h('Source index')
p('Paths are relative to the hippo repository unless explicitly identified as a report-local evidence file. The report-local evidence directory contains the retrieved aggregate records and a normalized machine-readable ledger. This index allows checking a number without treating an old narrative as the only evidence.')
sources=[
('S1','Consolidated low data ledger','experiments/consolidated_convergence_20260920/CONSOLIDATED.json','52 historical runs, selected and final audits, 75/200/400/600 budgets.'),
('S2','Fold zero edge comparison','docs/experiments/edge_lowdata_20260924/RESULTS.json','Three seeds, learning histories and paired train/validation audits.'),
('S3','Additional edge folds','docs/experiments/edge_lowdata_folds12_20260924/SUMMARY.json','Six completed runs; full error strata retrieved into evidence/edge_audits.json.'),
('S4','Fraction verification','evidence/contour_inventory.json','Contains the 10% cancelled-before-start submission receipt; 12.5% and 25% outcomes are in S1.'),
('S5','Early stopped augmentation replication','experiments/baseline_replication_20260916/three_seed_comparison/comparison.json','Three seed-matched full-data baseline and augmented comparisons.'),
('S6','Long schedule protocol','experiments/seed2_four_arm_600ep_20260919/PROTOCOL.md','Ten cases, 6,000 updates, no early stopping, StepLR interval 160.'),
('S7','Full data edge generalization audit','docs/experiments/edge_consistency_20260924/GENERALIZATION_RESULTS.json','Augmented Dice, bands and edge selected-checkpoint errors on train/validation.'),
('S8','Implemented edge and bands losses','thesis/new_constraints/edge_consistency.py and bands/outer_boundary.py','Exact six-face operator; the bands file is under thesis/new_constraints/.'),
('S9','Descending contour audit','docs/experiments/step_contour_20260922/README.md','260 reference geometries and matched early-stopped prediction analysis.'),
('S10','Sagittal step fitting','docs/experiments/step_contour_20260922/sagittal_step_fit.md','Isotonic function supervision and frozen-logit screens.'),
('S11','Withdrawn contour pilot','docs/experiments/step_contour_20260922/RESULTS.md','Retained historical summary; heavy artifacts were deleted before this review.'),
('S12','Independent frozen boundary head','docs/experiments/independent_surface_head_20260922/README.md','Head-only experiment, presence and endpoint disagreements.'),
('S13','Function feedback formulation','docs/experiments/step_contour_20260922/FUNCTION_HEAD_LOSS_DERIVATION_20260922.md','Probability anchors, supervised functions, responsible voxels and renderer.'),
('S14','Function robustness experiment','evidence/contour_inventory.json and evidence/contour_protocol.json','Retrieved original robustness summary and fold-1 source configuration.'),
('S15','Boundary preservation audit','docs/experiments/boundary_preservation_20260923/RESULTS.md','208-case matched error transitions; physical-distance strata.'),
('S16','Function component audit','experiments/presence_decisive_20260923/results_audit_666651/summary.json','Matched A, E raw, E final and labeled oracle interventions.'),
('S17','Minimum presence feasibility','docs/experiments/step_contour_20260922/PRESENCE_FEASIBILITY_RESULTS_20260923.md','Non-deployable oracle recovery limits.'),
('S18','Signed gradient diagnosis','docs/experiments/step_contour_20260922/SIGNED_GRADIENT_RESULTS_20260923.md','Output-level derivatives and shared-head directional responses.')]
for key,title,path,desc in sources:
    p(key+'  '+title,bold=True)
    q=p(path);q.paragraph_format.space_after=Pt(2)
    for run in q.runs:run.font.size=Pt(8)
    p(desc)

out=HERE/'Hippocampus_methods_and_results.docx'
doc.save(out)
print(out)
print('Tables',len(doc.tables),'paragraphs',len(doc.paragraphs),'master rows',sum(r['budget']==75 for r in R),'all low-data rows',len(R))
